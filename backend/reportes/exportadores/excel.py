"""Exportador Excel (openpyxl). Recibe el dict de contenido y devuelve bytes."""
import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .comun import NAVY, es_numero, fecha_generado, limpiar

RELLENO_NAVY = PatternFill('solid', start_color=NAVY, end_color=NAVY)
FUENTE_ENCABEZADO = Font(bold=True, color='FFFFFF')
_PROHIBIDOS = re.compile(r'[\[\]:*?/\\]')
MAX_ANCHO, MIN_ANCHO = 50, 10


def _nombre_hoja(titulo, usados):
    base = _PROHIBIDOS.sub(' ', limpiar(titulo)).strip().strip("'")[:31] or 'Hoja'
    nombre, n = base, 2
    while nombre.lower() in usados:
        sufijo = f' ({n})'
        nombre = base[:31 - len(sufijo)] + sufijo
        n += 1
    usados.add(nombre.lower())
    return nombre


def _escribir(hoja, fila, columna, valor, negrita=False):
    """Escribe una celda. Texto que empieza con = + - @ o tabulación se guarda como texto literal (sin fórmulas)."""
    celda = hoja.cell(row=fila, column=columna)
    if es_numero(valor):
        celda.value = valor
    else:
        texto = limpiar(valor) if valor is not None else ''
        celda.value = texto
        if texto and texto[0] in '=+-@\t':
            celda.data_type = 's'
    if negrita:
        celda.font = Font(bold=True)
    return celda


def _encabezados(hoja, fila, columnas):
    for j, nombre in enumerate(columnas, start=1):
        celda = _escribir(hoja, fila, j, nombre)
        celda.font = FUENTE_ENCABEZADO
        celda.fill = RELLENO_NAVY
        celda.alignment = Alignment(wrap_text=True, vertical='center')


def _ajustar_anchos(hoja):
    for columna in hoja.columns:
        mayor = max((len(str(c.value)) for c in columna if c.value is not None), default=0)
        hoja.column_dimensions[get_column_letter(columna[0].column)].width = min(max(mayor + 2, MIN_ANCHO), MAX_ANCHO)


def exportar(contenido):
    libro = Workbook()
    resumen = libro.active
    resumen.title = 'Resumen'
    usados = {'resumen'}

    _escribir(resumen, 1, 1, 'TRAINET — ' + str(contenido.get('titulo', '')), negrita=True).font = Font(bold=True, size=14, color=NAVY)
    _escribir(resumen, 2, 1, contenido.get('subtitulo', ''))
    _escribir(resumen, 3, 1, 'Generado el ' + fecha_generado(contenido))

    fila = 5
    if contenido.get('indicadores'):
        _encabezados(resumen, fila, ['Indicador', 'Valor'])
        for ind in contenido['indicadores']:
            fila += 1
            _escribir(resumen, fila, 1, ind['etiqueta'])
            _escribir(resumen, fila, 2, ind['valor'])
        fila += 2

    # Datos de los gráficos en la hoja Resumen y gráfico nativo de barras junto a ellos.
    for grafico in contenido.get('graficos', []):
        if not grafico['etiquetas']:
            continue
        _escribir(resumen, fila, 1, grafico['titulo'], negrita=True)
        fila += 1
        _encabezados(resumen, fila, ['Categoría'] + [s['nombre'] for s in grafico['series']])
        inicio = fila
        for i, etiqueta in enumerate(grafico['etiquetas']):
            fila += 1
            _escribir(resumen, fila, 1, etiqueta)
            for j, serie in enumerate(grafico['series'], start=2):
                _escribir(resumen, fila, j, serie['valores'][i])
        barras = BarChart()
        barras.type = 'col'
        barras.title = limpiar(grafico['titulo'])
        barras.add_data(Reference(resumen, min_col=2, max_col=1 + len(grafico['series']), min_row=inicio, max_row=fila),
                        titles_from_data=True)
        barras.set_categories(Reference(resumen, min_col=1, min_row=inicio + 1, max_row=fila))
        barras.width, barras.height = 18, 8
        resumen.add_chart(barras, f'F{inicio}')
        fila = max(fila, inicio + 17) + 2
    _ajustar_anchos(resumen)

    for seccion in contenido.get('secciones', []):
        hoja = libro.create_sheet(_nombre_hoja(seccion['titulo'], usados))
        _encabezados(hoja, 1, seccion['columnas'])
        for i, datos in enumerate(seccion['filas'], start=2):
            for j, valor in enumerate(datos, start=1):
                _escribir(hoja, i, j, valor)
        if seccion.get('nota'):
            _escribir(hoja, len(seccion['filas']) + 3, 1, seccion['nota'])
        hoja.freeze_panes = 'A2'
        _ajustar_anchos(hoja)

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()

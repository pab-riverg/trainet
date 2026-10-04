"""Exportador Word (python-docx). Recibe el dict de contenido y devuelve bytes."""
from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .comun import NAVY, columnas_maximas, cortar, fecha_generado, formatear

COLOR_NAVY = RGBColor.from_string(NAVY.upper())


def _sombrear(celda, color):
    propiedades = celda._tc.get_or_add_tcPr()
    relleno = OxmlElement('w:shd')
    relleno.set(qn('w:val'), 'clear')
    relleno.set(qn('w:color'), 'auto')
    relleno.set(qn('w:fill'), color)
    propiedades.append(relleno)


def _tabla(documento, columnas, filas):
    tabla = documento.add_table(rows=1, cols=len(columnas))
    tabla.style = 'Table Grid'
    for celda, nombre in zip(tabla.rows[0].cells, columnas):
        celda.text = ''
        corrida = celda.paragraphs[0].add_run(cortar(nombre))
        corrida.bold = True
        corrida.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        corrida.font.size = Pt(9)
        _sombrear(celda, NAVY)
    for datos in filas:
        celdas = tabla.add_row().cells
        for celda, valor in zip(celdas, datos):
            celda.text = ''
            celda.paragraphs[0].add_run(formatear(valor)).font.size = Pt(9)
    return tabla


def exportar(contenido):
    documento = Document()
    if columnas_maximas(contenido) > 6:
        seccion = documento.sections[0]
        seccion.orientation = WD_ORIENT.LANDSCAPE
        seccion.page_width, seccion.page_height = seccion.page_height, seccion.page_width
        seccion.left_margin = seccion.right_margin = Cm(1.5)

    titulo = documento.add_heading('TRAINET — ' + cortar(contenido.get('titulo', ''), 150), level=0)
    for corrida in titulo.runs:
        corrida.font.color.rgb = COLOR_NAVY
    documento.add_paragraph(cortar(contenido.get('subtitulo', '')))

    if contenido.get('indicadores'):
        documento.add_heading('Indicadores', level=1)
        _tabla(documento, ['Indicador', 'Valor'], [[i['etiqueta'], i['valor']] for i in contenido['indicadores']])

    for sec in contenido.get('secciones', []):
        documento.add_heading(cortar(sec['titulo']), level=1)
        if sec['filas']:
            _tabla(documento, sec['columnas'], sec['filas'])
        else:
            documento.add_paragraph('Sin datos.')
        if sec.get('nota'):
            documento.add_paragraph(cortar(sec['nota'], 300)).runs[0].italic = True

    # Los gráficos se describen como tabla de datos (sin imagen).
    for graf in contenido.get('graficos', []):
        if not graf['etiquetas']:
            continue
        documento.add_heading('Gráfico: ' + cortar(graf['titulo']), level=1)
        filas = [[etiqueta] + [s['valores'][i] for s in graf['series']] for i, etiqueta in enumerate(graf['etiquetas'])]
        _tabla(documento, ['Categoría'] + [s['nombre'] for s in graf['series']], filas)

    documento.add_paragraph('')
    nota = documento.add_paragraph('Informe generado por TRAINET el ' + fecha_generado(contenido) + '.')
    nota.runs[0].italic = True

    salida = BytesIO()
    documento.save(salida)
    return salida.getvalue()

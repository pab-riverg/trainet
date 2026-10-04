"""Lectura de archivos CSV/XLSX importados: validación, encabezados normalizados y filas de vista previa."""
import csv
import hashlib
import io
import json
import os
from datetime import date, datetime

MAX_BYTES = 10 * 1024 * 1024
MAX_FILAS = 50000
MAX_COLUMNAS = 50
EXTENSIONES = ('csv', 'xlsx')
DELIMITADORES = (',', ';', '\t')
CODIFICACIONES = ('utf-8-sig', 'latin-1')  # latin-1 acepta cualquier byte: es el último recurso


class ErrorImportacion(Exception):
    """Archivo no válido; el mensaje está en español y se muestra tal cual al usuario."""


def extension(nombre):
    return os.path.splitext(nombre)[1].lstrip('.').lower()


def huella(contenido):
    return hashlib.sha256(contenido).hexdigest()


def _texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, datetime):
        return valor.isoformat(sep=' ') if (valor.hour or valor.minute or valor.second) else valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _filas_csv(contenido):
    for codificacion in CODIFICACIONES:
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    primera = texto.split('\n', 1)[0]
    delimitador = max(DELIMITADORES, key=primera.count)
    if primera.count(delimitador) == 0:
        delimitador = ','
    for fila in csv.reader(io.StringIO(texto, newline=''), delimiter=delimitador):
        yield fila


def _filas_xlsx(contenido):
    from openpyxl import load_workbook
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise ErrorImportacion('El archivo no es un XLSX válido o está dañado.')
    try:
        hoja = libro.worksheets[0]
        for fila in hoja.iter_rows(values_only=True):
            yield list(fila)
    finally:
        libro.close()


def normalizar_encabezados(crudos):
    """Vacíos -> 'Columna N'; repetidos -> 'nombre_2', 'nombre_3'…"""
    resultado, usados = [], set()
    for i, valor in enumerate(crudos, start=1):
        nombre = _texto(valor) or f'Columna {i}'
        base, n = nombre, 2
        while nombre.lower() in usados:
            nombre = f'{base}_{n}'
            n += 1
        usados.add(nombre.lower())
        resultado.append(nombre)
    return resultado


def leer(contenido, formato, limite_filas=None):
    """Devuelve (encabezados, filas, total_filas). `filas` contiene hasta `limite_filas` filas (todas si es None).

    Ignora las filas totalmente vacías. Lanza ErrorImportacion si el contenido no cumple los requisitos.
    """
    origen = _filas_csv(contenido) if formato == 'csv' else _filas_xlsx(contenido)
    encabezados, filas, total = None, [], 0
    for cruda in origen:
        if encabezados is None:
            if not any(_texto(c) for c in cruda):
                continue  # líneas en blanco antes del encabezado
            ancho = len(cruda)
            while ancho and not _texto(cruda[ancho - 1]):
                ancho -= 1  # columnas vacías al final
            if ancho > MAX_COLUMNAS:
                raise ErrorImportacion(f'El archivo tiene {ancho} columnas; el máximo permitido es {MAX_COLUMNAS}.')
            encabezados = normalizar_encabezados(cruda[:ancho])
            continue
        valores = [_texto(c) for c in cruda[:len(encabezados)]]
        if not any(valores):
            continue
        total += 1
        if total > MAX_FILAS:
            raise ErrorImportacion(f'El archivo supera el máximo de {MAX_FILAS} filas de datos.')
        valores += [''] * (len(encabezados) - len(valores))
        if limite_filas is None or len(filas) < limite_filas:
            filas.append(valores)
    if encabezados is None:
        raise ErrorImportacion('El archivo está vacío.')
    if total == 0:
        raise ErrorImportacion('El archivo solo tiene encabezados; debe incluir al menos una fila de datos.')
    return encabezados, filas, total


def analizar(nombre, contenido):
    """Valida un archivo subido y devuelve {formato, encabezados, filas, columnas, huella}."""
    formato = extension(nombre)
    if formato not in EXTENSIONES:
        raise ErrorImportacion('Tipo de archivo no permitido. Solo se aceptan archivos .csv y .xlsx.')
    if len(contenido) > MAX_BYTES:
        raise ErrorImportacion('El archivo supera el límite de 10 MB.')
    if not contenido:
        raise ErrorImportacion('El archivo está vacío.')
    encabezados, _, total = leer(contenido, formato, limite_filas=0)
    return {
        'formato': formato,
        'encabezados': json.dumps(encabezados, ensure_ascii=False),
        'filas': total,
        'columnas': len(encabezados),
        'huella': huella(contenido),
    }
